import copy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DESIGN = "充电桩车位防占系统_设计文档.docx"
OUTLINE = "充电桩车位防占系统_开发大纲.docx"

def all_paragraphs(doc):
    paras = list(doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                paras.extend(cell.paragraphs)
    return paras

def replace_text(doc, repls):
    counts = {}
    for p in all_paragraphs(doc):
        for run in p.runs:
            for old, new in repls:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    counts[old] = counts.get(old, 0) + 1
    return counts

def find_table(doc, marker):
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                if marker in cell.text:
                    return tbl
    return None

def add_rows(doc, marker, rows):
    tbl = find_table(doc, marker)
    if tbl is None:
        return False
    for cells_text in rows:
        r = tbl.add_row()
        for i, txt in enumerate(cells_text):
            r.cells[i].text = txt
    return True

def insert_bullet_after(doc, anchor_text, texts):
    target = None
    for p in doc.paragraphs:
        if anchor_text in p.text:
            target = p
            break
    if target is None:
        return False
    current = target._p
    ref_runs = target._p.findall(qn('w:r'))
    rpr = ref_runs[0].find(qn('w:rPr')) if ref_runs else None
    for txt in texts:
        new_el = copy.deepcopy(target._p)
        for r in new_el.findall(qn('w:r')):
            new_el.remove(r)
        run = OxmlElement('w:r')
        if rpr is not None:
            run.append(copy.deepcopy(rpr))
        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = txt
        run.append(t)
        new_el.append(run)
        current.addnext(new_el)
        current = new_el
    return True

# ---- 设计文档 ----
d = Document(DESIGN)
repls_d = [
    ("PaddleOCR/HyperLPR", "HyperLPR（主）/ PaddleOCR（备选）"),
    ("PaddleOCR / HyperLPR", "HyperLPR（主）/ PaddleOCR（备选）"),
    ("Vue 3 + Vite", "Vue 3 + Vite + Element Plus"),
    ("MVP 可先用极简单页", "MVP 极简后台；预览仅取最近一帧"),
]
print("design repl:", replace_text(d, repls_d))
insert_bullet_after(d, "视频源：开发期用测试视频", [
    "部署形态：推理（抽帧→检测→识牌→规则）与 Web 服务拆分为独立进程，Web 仅读库并供前端。",
    "前端预览：仅展示最近一帧截图（轮询刷新），不接实时视频流。",
])
print("design add_rows(数据库):", add_rows(d, "数据库", [
    ["部署", "推理与 Web 服务拆分为独立进程", "Web 仅读库 + 前端接口"],
    ["预览", "最近一帧截图轮询", "不接实时视频流"],
]))
d.save(DESIGN)

# ---- 开发大纲 ----
o = Document(OUTLINE)
repls_o = [
    ("PaddleOCR / HyperLPR", "HyperLPR（主）/ PaddleOCR（备选）"),
    ("PaddleOCR/HyperLPR", "HyperLPR（主）/ PaddleOCR（备选）"),
]
print("outline repl:", replace_text(o, repls_o))
print("outline add_rows(车牌识别):", add_rows(o, "车牌识别", [
    ["部署架构", "推理进程与 Web 服务拆分（独立进程；Web 仅读库 + 供前端）"],
    ["前端预览", "仅取最近一帧截图（轮询刷新），不接实时视频流"],
]))
o.save(OUTLINE)

# ---- 校验 ----
for path in [DESIGN, OUTLINE]:
    dd = Document(path)
    print("=== VERIFY", path, "===")
    for p in all_paragraphs(dd):
        t = p.text.strip()
        if any(k in t for k in ["HyperLPR", "独立进程", "最近一帧", "Element Plus",
                                 "部署形态", "前端预览", "部署架构", "PaddleOCR"]):
            print(" -", t)
print("DONE")
